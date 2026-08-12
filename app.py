"""
FastAPI Server for PIRL DOL Error Report Inspector
Provides REST API endpoints and web interface for uploading PIRL data,
parsing DOL/WIPS error logs or copy-pasted text, searching reconstructed PIRL tables,
cell editing, and headerless CSV exports.
"""

import os
import io
from datetime import datetime
import pandas as pd
from fastapi import FastAPI, UploadFile, File, Query, Request, HTTPException, Form, Body
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from typing import Optional, List
from pirl_parser import PIRLParser

app = FastAPI(title="PIRL DOL Error Report Inspector", version="1.0.0")

# Setup templates and static files directories
os.makedirs("templates", exist_ok=True)
os.makedirs("static", exist_ok=True)

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

DEFAULT_SCHEMA_PATH = "PIRL Schema Header.csv" if os.path.exists("PIRL Schema Header.csv") else None
parser = PIRLParser(schema_header_path=DEFAULT_SCHEMA_PATH)

class PasteErrorRequest(BaseModel):
    raw_text: str

class CellUpdateRequest(BaseModel):
    pirl_row_index: int
    col_name_or_index: str
    new_value: str

@app.get("/", response_class=HTMLResponse)
async def serve_gui(request: Request):
    """Renders the main PIRL DOL Error Viewer dashboard."""
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "schema_loaded": len(parser.schema_columns) > 0,
            "schema_column_count": len(parser.schema_columns)
        }
    )

@app.get("/api/schema")
async def get_schema_info():
    """Returns PIRL schema columns and element metadata."""
    return {
        "loaded": len(parser.schema_columns) > 0,
        "count": len(parser.schema_columns),
        "columns": parser.schema_columns
    }

@app.post("/api/upload")
async def upload_files(
    pirl_file: Optional[UploadFile] = File(None),
    error_file: Optional[UploadFile] = File(None),
    schema_file: Optional[UploadFile] = File(None)
):
    """Processes uploaded PIRL Data, DOL Error Reports, and/or PIRL Schema Header files."""
    results = {}

    if schema_file and schema_file.filename:
        schema_bytes = await schema_file.read()
        cols = parser.load_schema_header(io.BytesIO(schema_bytes))
        results["schema_columns_count"] = len(cols)
    elif not parser.schema_columns and os.path.exists("PIRL Schema Header.csv"):
        parser.load_schema_header("PIRL Schema Header.csv")

    if pirl_file and pirl_file.filename:
        pirl_bytes = await pirl_file.read()
        records_count = parser.parse_pirl_data(pirl_bytes, pirl_file.filename)
        results["pirl_records_processed"] = records_count
        results["pirl_filename"] = pirl_file.filename

    if error_file and error_file.filename:
        err_bytes = await error_file.read()
        parsed_errors = parser.parse_error_report(err_bytes, error_file.filename)
        results["errors_parsed"] = len(parsed_errors)
        results["error_filename"] = error_file.filename

    results["stats"] = parser.get_stats()
    return JSONResponse(content=results)

@app.post("/api/paste-errors")
async def paste_errors(payload: PasteErrorRequest):
    """Processes raw text copy-pasted directly from WIPS edit check results web page."""
    parsed_errors = parser.parse_pasted_wips_text(payload.raw_text)
    return {
        "status": "success",
        "errors_parsed": len(parsed_errors),
        "stats": parser.get_stats()
    }

@app.post("/api/preflight-check")
async def run_preflight_check():
    """Runs preflight logical & WIPS rule validation check against current PIRL dataset."""
    if parser.pirl_df is None or parser.pirl_df.empty:
        raise HTTPException(status_code=400, detail="No PIRL dataset loaded to run preflight check.")

    pre_discrepancies = parser.run_pirl_pre_validation_checks()
    return {
        "status": "success",
        "preflight_discrepancies_count": len(pre_discrepancies),
        "stats": parser.get_stats(),
        "message": f"Preflight check completed! Detected {len(pre_discrepancies)} logical discrepancy item(s)."
    }

@app.get("/api/pirl-data")
async def get_pirl_table_data(
    query: str = Query("", description="Global text search"),
    uiid: str = Query("", description="Search by Unique ID (Element 100)"),
    state_id: str = Query("", description="Search by State ID (Element 101)"),
    dob: str = Query("", description="Search by Date of Birth (Element 200)"),
    veteran: str = Query("", description="Search by Veteran Status (Element 300)"),
    disability: str = Query("", description="Search by Disability (Element 202)"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200)
):
    """Returns paginated reconstructed PIRL table rows with column filters."""
    return parser.get_pirl_table_data(
        query=query,
        uiid=uiid,
        state_id=state_id,
        dob=dob,
        veteran=veteran,
        disability=disability,
        page=page,
        page_size=page_size
    )

@app.get("/api/errors")
async def get_errors(
    query: str = Query("", description="Search term across UIID, element, code, message"),
    element: str = Query("", description="Filter by PIRL element number/name"),
    severity: str = Query("", description="Filter by Error, Warning, or Info"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200)
):
    """Returns paginated and filtered list of errors."""
    return parser.get_filtered_errors(query=query, element_filter=element, severity_filter=severity, page=page, page_size=page_size)

@app.get("/api/stats")
async def get_stats():
    """Returns overall statistical metrics."""
    return parser.get_stats()

@app.get("/api/record/{row_index}")
async def get_record_details(row_index: int):
    """Returns complete PIRL row field values with failing field indicators."""
    details = parser.get_full_record_details(row_index)
    if details is None:
        raise HTTPException(status_code=404, detail="PIRL Record index out of range or no PIRL data loaded.")
    return details

@app.put("/api/cell-update")
async def update_cell(payload: CellUpdateRequest):
    """Updates a cell value in the PIRL dataframe in memory and saves directly back to file in folder."""
    success = parser.update_cell_value(payload.pirl_row_index, payload.col_name_or_index, payload.new_value)
    if not success:
        raise HTTPException(status_code=400, detail="Could not update cell value. Check row index or column name.")
    
    saved_disk, msg = parser.save_to_disk()
    return {
        "status": "success",
        "row_index": payload.pirl_row_index,
        "col": payload.col_name_or_index,
        "new_value": payload.new_value,
        "disk_saved": saved_disk,
        "save_message": msg,
        "loaded_file": parser.loaded_filepath or parser.loaded_filename,
        "stats": parser.get_stats()
    }

@app.post("/api/save-file")
async def save_file_to_disk():
    """Saves current corrected dataset directly back to original file on disk without changing CSV format."""
    saved_disk, msg = parser.save_to_disk()
    return {
        "status": "success" if saved_disk else "error",
        "message": msg,
        "loaded_file": parser.loaded_filepath or parser.loaded_filename,
        "stats": parser.get_stats()
    }

@app.get("/api/export-headerless-pirl")
async def export_headerless_pirl():
    """Exports the modified PIRL dataset as a headerless CSV file formatted for WIPS upload (wp_pirl_{date}.csv)."""
    csv_bytes, filename = parser.generate_headerless_pirl_csv()
    return StreamingResponse(
        io.BytesIO(csv_bytes),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@app.post("/api/demo")
async def load_demo():
    """Loads instant sample PIRL data and sample DOL errors for testing."""
    if not parser.schema_columns and os.path.exists("PIRL Schema Header.csv"):
        parser.load_schema_header("PIRL Schema Header.csv")
    
    rows, errs = parser.generate_demo_data()
    return {
        "status": "success",
        "demo_rows": rows,
        "demo_errors": errs,
        "stats": parser.get_stats()
    }

@app.get("/api/export-errors")
async def export_errors(
    query: str = Query(""),
    element: str = Query(""),
    severity: str = Query(""),
    format: str = Query("csv", pattern="^(csv|xlsx)$")
):
    """Exports filtered error list with participant details as CSV or Excel file download."""
    result = parser.get_filtered_errors(query=query, element_filter=element, severity_filter=severity, page=1, page_size=100000)
    errs = result["errors"]

    export_data = []
    for e in errs:
        school_st = e.get("school_status", "")
        school_label = "5 - Sec. Grad" if school_st == "5" else ("1 - Attending" if school_st == "1" else f"Code {school_st}" if school_st else "")

        export_data.append({
            "PIRL Row #": e.get("row_number", ""),
            "Unique ID (Elem 100)": e.get("uiid", ""),
            "PIRL Element": e.get("element_col_name", e.get("element", "")),
            "Element #": e.get("element_number", ""),
            "Current Value": e.get("current_value", ""),
            "Program Entry Date (900)": e.get("entry_date", ""),
            "Program Exit Date (901)": e.get("exit_date", ""),
            "School Status (409)": school_label,
            "State Code (101)": e.get("state_code", ""),
            "Error Code": e.get("error_code", ""),
            "Error Description": e.get("error_message", ""),
            "Severity": e.get("severity", "Error")
        })

    df = pd.DataFrame(export_data)

    if format == "xlsx":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PIRL Error Audit & Resolution"
        ws.views.sheetView[0].showGridLines = True

        # Title Block
        ws.merge_cells("A1:M1")
        t_cell = ws["A1"]
        t_cell.value = "PIRL WIPS Edit Check Error Audit & Resolution Report"
        t_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:M2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Generated Audit Report | Total Errors: {len(errs)} | Grouped by Unique Individual Identifier (Elem 100)"
        sub_cell.font = Font(name="Segoe UI", size=10, italic=True, color="475569")
        ws.row_dimensions[2].height = 20

        # Table Headers
        headers = [
            "PIRL Row #",
            "Unique ID (Elem 100)",
            "PIRL Element",
            "Element #",
            "Current Value",
            "Program Entry Date (900)",
            "Program Exit Date (901)",
            "School Status (409)",
            "State Code (101)",
            "Error Code",
            "Error Description",
            "Severity",
            "Resolution / Audit Notes"
        ]

        h_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        h_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        h_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[4].height = 28
        for c_idx, h in enumerate(headers, 1):
            c = ws.cell(row=4, column=c_idx)
            c.value = h
            c.fill = h_fill
            c.font = h_font
            c.alignment = h_align

        # Group Fills and Cell Styles
        group_fills = [
            PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"),
            PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid"),
        ]
        val_err_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        val_err_font = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

        notes_fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
        notes_font = Font(name="Segoe UI", size=10, italic=True, color="166534")

        thin_b = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )
        group_top_b = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="medium", color="2563EB"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        c_row = 5
        prev_uiid = None
        fill_toggle = 0

        for e in errs:
            uiid = str(e.get("uiid", "")).strip()
            if prev_uiid is not None and uiid != prev_uiid:
                fill_toggle = (fill_toggle + 1) % 2

            r_fill = group_fills[fill_toggle]
            r_border = group_top_b if (prev_uiid is not None and uiid != prev_uiid) else thin_b

            school_st = e.get("school_status", "")
            school_label = "5 - Sec. Grad" if school_st == "5" else ("1 - Attending" if school_st == "1" else f"Code {school_st}" if school_st else "")

            row_vals = [
                e.get("row_number", ""),
                uiid,
                e.get("element_col_name", e.get("element", "")),
                e.get("element_number", ""),
                e.get("current_value", ""),
                e.get("entry_date", ""),
                e.get("exit_date", ""),
                school_label,
                e.get("state_code", ""),
                e.get("error_code", ""),
                e.get("error_message", ""),
                e.get("severity", "Error"),
                "" # Blank editable column for Resolution / Audit Notes
            ]

            ws.row_dimensions[c_row].height = 22
            for col_n, val in enumerate(row_vals, 1):
                cell = ws.cell(row=c_row, column=col_n)
                cell.value = val
                cell.fill = r_fill
                cell.border = r_border
                cell.font = Font(name="Segoe UI", size=10)

                if col_n in [1, 2, 4, 5, 6, 7, 9, 10, 12]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

                if col_n == 5 and e.get("element_number") == "200":
                    cell.fill = val_err_fill
                    cell.font = val_err_font

                if col_n == 2:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="1E40AF")

                if col_n == 13:
                    cell.fill = notes_fill
                    cell.font = notes_font
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            prev_uiid = uiid
            c_row += 1

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.row in [1, 2]: continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        ws.column_dimensions["K"].width = 45
        ws.column_dimensions["M"].width = 38

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=pirl_wips_errors_report.xlsx"}
        )
    else:
        # Include editable Resolution / Audit Notes column in CSV export too
        for d in export_data:
            d["Resolution / Audit Notes"] = ""
        df = pd.DataFrame(export_data)
        output = io.StringIO()
        df.to_csv(output, index=False)
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode('utf-8')),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=pirl_wips_errors_report.csv"}
        )

@app.get("/api/export-pirl-headers")
async def export_pirl_headers(format: str = Query("csv", pattern="^(csv|xlsx)$")):
    """Exports the reconstructed PIRL dataset with full 495 column headers as CSV or Excel file download."""
    if parser.pirl_df is None or parser.pirl_df.empty:
        raise HTTPException(status_code=400, detail="No PIRL data loaded to export.")

    df_export = parser.pirl_df.copy()
    if parser.schema_columns and len(parser.schema_columns) == len(df_export.columns):
        df_export.columns = parser.schema_columns

    date_str = datetime.now().strftime("%Y%m%d")

    if format == "xlsx":
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_export.head(50000).to_excel(writer, index=False, sheet_name="PIRL Reconstructed")
        output.seek(0)
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename=pirl_full_reconstructed_{date_str}.xlsx"}
        )
    else:
        output = io.StringIO()
        df_export.to_csv(output, index=False)
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode('utf-8')),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename=pirl_full_reconstructed_{date_str}.csv"}
        )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=True)
